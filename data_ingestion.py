import os
import re
import json
import uuid
import pdfplumber
import tiktoken
from pathlib import Path
from typing import Callable, List, Dict, Any

try:
    import docx  # python-docx
    _DOCX_AVAILABLE = True
except ImportError:
    _DOCX_AVAILABLE = False

def extract_text_from_txt(file_path: str) -> str:
    """Reads raw text from a .txt file."""
    with open(file_path, "r", encoding="utf-8") as f:
        return f.read()

def extract_text_from_pdf(file_path: str) -> str:
    """Extracts text from all pages of a .pdf file using pdfplumber."""
    text = ""
    try:
        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                # Some pages might be scanned images with no text
                if page_text:
                    text += page_text + "\n"
    except Exception as e:
        print(f"Error reading PDF {file_path}: {e}")
    return text


def extract_text_from_docx(file_path: str) -> str:
    """Extracts text from a .docx file using python-docx.

    Joins paragraph text with newlines, skipping empty paragraphs.
    Tables are not extracted (body text only) — extend if needed.

    Raises:
        ImportError  -- if python-docx is not installed.
        RuntimeError -- if the file is corrupted or not a valid DOCX.
    """
    if not _DOCX_AVAILABLE:
        raise ImportError(
            "python-docx is required to process DOCX files. "
            "Install it with: pip install python-docx"
        )
    try:
        document = docx.Document(file_path)
        paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
        return "\n".join(paragraphs)
    except Exception as e:
        raise RuntimeError(f"Failed to read DOCX file '{file_path}': {e}") from e

def chunk_text(text: str, chunk_size: int = 500, overlap: int = 50) -> List[str]:
    """
    Splits text into chunks of roughly `chunk_size` tokens,
    with an overlap of `overlap` tokens between consecutive chunks.
    """
    # We use 'tiktoken' to count actual tokens rather than words.
    # 'cl100k_base' is the standard tokenizer used by modern OpenAI models (like GPT-4),
    # which is a great baseline for general RAG systems.
    encoding = tiktoken.get_encoding("cl100k_base")
    tokens = encoding.encode(text)
    
    chunks = []
    
    # If the text is empty, return an empty list
    if len(tokens) == 0:
        return []
    
    # How the overlap logic works:
    # We iterate over the list of tokens. Instead of jumping forward by the full 
    # `chunk_size` (which would give zero overlap), we jump forward by `chunk_size - overlap`.
    # This means the last `overlap` number of tokens from the current chunk will be included 
    # at the start of the next chunk. 
    # Overlap is crucial in RAG so we don't accidentally split a sentence or concept in half!
    step = chunk_size - overlap
    
    for i in range(0, len(tokens), step):
        # Slice the tokens list from the current index `i` to `i + chunk_size`
        chunk_tokens = tokens[i:i + chunk_size]
        
        # Decode the tokens back into a readable string
        chunk_text = encoding.decode(chunk_tokens)
        chunks.append(chunk_text)
        
    return chunks

def process_documents(input_dir: str, output_file: str, chunk_size: int = 500, overlap: int = 50):
    """
    Loads documents from a directory, extracts text, chunks it, and saves to JSON.
    """
    input_path = Path(input_dir)
    if not input_path.exists() or not input_path.is_dir():
        print(f"Directory not found: {input_dir}")
        return

    all_chunks: List[Dict[str, Any]] = []
    
    # Iterate through all files in the given directory
    for file_name in os.listdir(input_dir):
        file_path = input_path / file_name
        
        # Skip if it's a directory
        if not file_path.is_file():
            continue
            
        ext = file_path.suffix.lower()
        
        # Extract text based on file type
        if ext == '.txt':
            print(f"Processing TXT: {file_name}")
            text = extract_text_from_txt(str(file_path))
        elif ext == '.pdf':
            print(f"Processing PDF: {file_name}")
            text = extract_text_from_pdf(str(file_path))
        elif ext == '.docx':
            print(f"Processing DOCX: {file_name}")
            try:
                text = extract_text_from_docx(str(file_path))
            except ImportError as e:
                print(f"Error reading DOCX {file_name}: {e}")
                text = ""
            # RuntimeError (corrupt/invalid file) is intentionally NOT caught here
            # — it propagates up to api.py's pipeline handler, which converts it
            # to a clean HTTP 500 with a meaningful message.
        else:
            # Skip truly unsupported formats (images, spreadsheets, etc.)
            print(f"Skipping unsupported format: {file_name}")
            continue
            
        # Ignore empty documents
        if not text.strip():
            print(f"Warning: No text found in {file_name}")
            continue
            
        # Split text into tokenized chunks
        text_chunks = chunk_text(text, chunk_size, overlap)
        
        # Store metadata and chunk info
        for index, chunk in enumerate(text_chunks):
            chunk_data = {
                "chunk_id": f"{file_name}_chunk_{index}",  # Unique deterministic ID for each chunk
                "source_file": file_name,       # Keep track of where it came from
                "text": chunk,                  # The actual raw text of the chunk
                "chunk_index": index            # Its position in the source document
            }
            all_chunks.append(chunk_data)
            
    # Save the results to a JSON file
    with open(output_file, 'w', encoding='utf-8') as f:
        # ensure_ascii=False handles special characters properly
        json.dump(all_chunks, f, indent=4, ensure_ascii=False)
        
    print(f"Successfully processed {len(all_chunks)} chunks from {input_dir}.")
    print(f"Output saved to {output_file}")


# ---------------------------------------------------------------------------
# Structure-aware extraction and chunking
# ---------------------------------------------------------------------------
# The functions above (extract_text_from_* + chunk_text) are the original
# pipeline: pdfplumber text cut into fixed 500-token tiktoken windows.  The
# gold-set evaluation (eval/) showed three problems with them:
#   - pdfplumber glues words together on many PDFs ("struggletomemorise...")
#     and interleaves the lines of two-column layouts;
#   - a third of the 500 cl100k-token windows exceed bge-small's 512-token
#     limit, so their tails are silently never embedded;
#   - fixed windows cut sentences, table rows and sections in half.
# The functions below replace them:
#   extract_pages()  - PDFium for PDFs (correct reading order and spacing),
#                      DOCX paragraphs AND tables in document order, page
#                      numbers kept for citations.
#   chunk_pages()    - packs whole sentences up to a budget measured with the
#                      embedder's own tokenizer, starts a new chunk at section
#                      headings, overlaps by whole sentences, and labels each
#                      chunk with its section heading.

_PAGE_NUMBER_LINE = re.compile(r"^\s*\d{1,4}\s*$")
# "3.2 Model architectures", "A.1.4. Pre-Deployment Testing", "CHAPTER 2" style headings
_HEADING = re.compile(
    r"^((\d+(\.\d+)*\.?|[A-Z]\.\d+(\.\d+)*\.?)\s+[A-Z][^.!?]{1,80}|CHAPTER\s+\d+.{0,60}|APPENDIX\s+[A-Z].{0,60})$"
)
_ENDS_SENTENCE = re.compile(r"[.!?:;\"”)]\s*$")
_SENTENCE_SPLIT = re.compile(r"(?<=[.!?])\s+(?=[A-Z0-9\"“(])")


def _unit(text: str, heading: bool = False) -> dict:
    return {"text": text, "heading": heading}


def _pdf_pages(file_path: str) -> list[dict]:
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(file_path)
    try:
        pages = []
        for i in range(len(pdf)):
            text = pdf[i].get_textpage().get_text_range()
            # PDFium marks line-break hyphenation with U+FFFE: rejoin the split word.
            text = text.replace("\r\n", "\n").replace("\r", "\n").replace("￾\n", "").replace("￾", "")
            units = []
            for line in text.split("\n"):
                line = line.strip()
                if line and not _PAGE_NUMBER_LINE.match(line):
                    units.append(_unit(line, bool(_HEADING.match(line))))
            pages.append({"page": i + 1, "units": units})
        return pages
    finally:
        pdf.close()


def _docx_pages(file_path: str) -> list[dict]:
    if not _DOCX_AVAILABLE:
        raise ImportError(
            "python-docx is required to process DOCX files. "
            "Install it with: pip install python-docx"
        )
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        document = docx.Document(file_path)
    except Exception as e:
        raise RuntimeError(f"Failed to read DOCX file '{file_path}': {e}") from e

    units = []
    # Walk the body in document order so tables stay next to the text that introduces them.
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(child, document)
            text = " ".join(para.text.split())
            if text:
                style = para.style.name if para.style is not None else ""
                units.append(_unit(text, style.startswith(("Heading", "Title")) or bool(_HEADING.match(text))))
        elif tag == "tbl":
            for row in Table(child, document).rows:
                cells: list[str] = []
                for cell in row.cells:
                    cell_text = " ".join(cell.text.split())
                    # Horizontally merged cells repeat their text once per grid column.
                    if cell_text and (not cells or cells[-1] != cell_text):
                        cells.append(cell_text)
                if cells:
                    units.append(_unit(" | ".join(cells)))
    return [{"page": None, "units": units}]


def extract_pages(file_path: str) -> list[dict]:
    """
    Extracts a document as a list of pages, each {"page": int | None, "units": [...]}.
    A unit is one line (PDF/TXT), paragraph or table row (DOCX), flagged when it
    looks like a section heading.  DOCX has no fixed pages, so page is None.

    Raises:
        ValueError   -- unsupported extension.
        ImportError  -- python-docx missing for a .docx file.
        RuntimeError -- corrupt DOCX.
    """
    ext = Path(file_path).suffix.lower()
    if ext == ".pdf":
        return _pdf_pages(file_path)
    if ext == ".docx":
        return _docx_pages(file_path)
    if ext == ".txt":
        units = [_unit(line.strip(), bool(_HEADING.match(line.strip())))
                 for line in extract_text_from_txt(file_path).splitlines() if line.strip()]
        return [{"page": None, "units": units}]
    raise ValueError(f"Unsupported file type: {ext}")


def _split_oversized(text: str, count_tokens: Callable[[str], int], max_tokens: int) -> list[str]:
    """Splits a unit longer than max_tokens into sentence-packed pieces (words as a last resort)."""
    if count_tokens(text) <= max_tokens:
        return [text]
    pieces: list[str] = []
    current = ""
    for sentence in _SENTENCE_SPLIT.split(text):
        if count_tokens(sentence) > max_tokens:
            words = sentence.split()
            sentence_parts, part = [], []
            for word in words:
                if part and count_tokens(" ".join(part + [word])) > max_tokens:
                    sentence_parts.append(" ".join(part))
                    part = []
                part.append(word)
            sentence_parts.append(" ".join(part))
        else:
            sentence_parts = [sentence]
        for s in sentence_parts:
            candidate = f"{current} {s}".strip()
            if current and count_tokens(candidate) > max_tokens:
                pieces.append(current)
                current = s
            else:
                current = candidate
    if current:
        pieces.append(current)
    return pieces


def chunk_pages(
    pages: list[dict],
    count_tokens: Callable[[str], int],
    max_tokens: int = 300,
    overlap_tokens: int = 50,
) -> list[dict]:
    """
    Packs extracted units into chunks of at most ~max_tokens embedder tokens.

    1. Units are grouped into sentence-complete segments: consecutive lines are
       joined until one ends a sentence, so a chunk boundary never falls inside
       a sentence that wraps across lines.  Headings are their own segment.
    2. Segments are packed greedily.  A heading starts a new chunk once the
       current one is at least 40% full, so sections aren't split mid-way.
    3. Consecutive chunks overlap by whole trailing segments (up to
       overlap_tokens), never across a heading.

    Returns dicts with: text (section heading prepended), page (first page),
    pages (all pages spanned), section.
    """
    # -- 1. sentence-complete segments --------------------------------------
    segments: list[dict] = []
    seg_units: list[dict] = []

    def close_segment():
        nonlocal seg_units
        if seg_units:
            segments.append({
                "units": seg_units,
                "tokens": sum(u["tokens"] for u in seg_units),
                "heading": seg_units[0]["heading"],
            })
            seg_units = []

    section = None
    for page in pages:
        for raw in page["units"]:
            if raw["heading"]:
                close_segment()
                section = raw["text"]
            for piece in _split_oversized(raw["text"], count_tokens, max_tokens):
                unit = {
                    "text": piece,
                    "page": page["page"],
                    "section": section,
                    "heading": raw["heading"],
                    "tokens": count_tokens(piece) + 1,  # +1 for the newline joiner
                }
                if seg_units and sum(u["tokens"] for u in seg_units) + unit["tokens"] > max_tokens:
                    close_segment()
                seg_units.append(unit)
                if raw["heading"] or _ENDS_SENTENCE.search(piece):
                    close_segment()
    close_segment()

    # -- 2/3. pack segments into overlapping chunks --------------------------
    chunks: list[dict] = []

    def emit(segs: list[dict]):
        units = [u for s in segs for u in s["units"]]
        if not units:
            return
        body = "\n".join(u["text"] for u in units)
        sec = units[0]["section"]
        text = body if not sec or body.startswith(sec) else f"{sec}\n{body}"
        page_numbers = sorted({u["page"] for u in units if u["page"] is not None})
        chunks.append({
            "text": text,
            "page": page_numbers[0] if page_numbers else None,
            "pages": page_numbers,
            "section": sec,
        })

    current: list[dict] = []
    for seg in segments:
        size = sum(s["tokens"] for s in current)
        if current and (size + seg["tokens"] > max_tokens or (seg["heading"] and size >= 0.4 * max_tokens)):
            emit(current)
            overlap: list[dict] = []
            if not seg["heading"]:
                for prev in reversed(current):
                    if sum(s["tokens"] for s in overlap) + prev["tokens"] > overlap_tokens:
                        break
                    overlap.insert(0, prev)
            current = overlap if sum(s["tokens"] for s in overlap) + seg["tokens"] <= max_tokens else []
        current.append(seg)
    emit(current)
    return chunks


def _estimate_tokens(text: str) -> int:
    """Fallback when the embedder's tokenizer isn't available: ~1.4 WordPiece tokens per word."""
    return int(len(text.split()) * 1.4) + 1


def ingest_documents(
    input_dir: str,
    output_file: str,
    count_tokens: Callable[[str], int] | None = None,
    max_tokens: int = 300,
    overlap_tokens: int = 50,
) -> list[dict]:
    """
    Structure-aware replacement for process_documents(): extracts every
    supported file in input_dir, chunks it with chunk_pages(), writes the
    chunks to output_file and returns them.

    count_tokens should count tokens the way the embedding model does
    (vector_store.embedding_token_counter()); a word-based estimate is used
    otherwise.
    """
    count_tokens = count_tokens or _estimate_tokens
    all_chunks: list[dict] = []
    for file_path in sorted(Path(input_dir).iterdir()):
        if not file_path.is_file():
            continue
        if file_path.suffix.lower() not in (".pdf", ".docx", ".txt"):
            print(f"Skipping unsupported format: {file_path.name}")
            continue
        # RuntimeError (corrupt DOCX) propagates to api.py, which returns a clean HTTP 500.
        pages = extract_pages(str(file_path))
        chunks = chunk_pages(pages, count_tokens, max_tokens, overlap_tokens)
        if not chunks:
            print(f"Warning: No text found in {file_path.name}")
            continue
        for index, chunk in enumerate(chunks):
            all_chunks.append({
                "chunk_id": f"{file_path.name}_chunk_{index}",
                "source_file": file_path.name,
                "chunk_index": index,
                **chunk,
            })
        print(f"Processed {file_path.name}: {len(chunks)} chunks")

    with open(output_file, "w", encoding="utf-8") as f:
        json.dump(all_chunks, f, indent=4, ensure_ascii=False)
    return all_chunks


if __name__ == "__main__":
    # Settings for our RAG ingestion
    DOCS_FOLDER = "docs"
    OUTPUT_JSON = "processed_chunks.json"
    CHUNK_SIZE = 150
    OVERLAP_SIZE = 30
    
    # Create the docs folder if it doesn't exist so you can easily drop files in it
    os.makedirs(DOCS_FOLDER, exist_ok=True)
    
    print(f"Looking for documents in '{DOCS_FOLDER}'...")
    process_documents(
        input_dir=DOCS_FOLDER, 
        output_file=OUTPUT_JSON, 
        chunk_size=CHUNK_SIZE, 
        overlap=OVERLAP_SIZE
    )
